"""Strict, privacy-safe FamilyBiz XLSX inspection and normalization."""

from __future__ import annotations

import hashlib
import io
import json
import re
import unicodedata
import zipfile
from collections import Counter, deque
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string
from openpyxl.utils.exceptions import InvalidFileException

from family_finance.config import Settings
from family_finance.models import (
    AccountDescriptor,
    DataQualityIssue,
    ImportInspection,
    ParsedSourceRecord,
)

HEADERS = [
    "תאריך",
    "סכום",
    "תאור",
    "תאריך ביעדים",
    "סוג תנועה",
    "קטגוריה",
    "מטבע",
    "מטבע מקורי",
    "תנועה מקורית",
]
_RAW_PAYLOAD_KEYS = tuple(str(index) for index in range(1, 11))
REPORT_END_RE = re.compile(r"(\d{2}/\d{2}/\d{4})\s*:עד תאריך")
REPORT_START_RE = re.compile(r"(\d{2}/\d{2}/\d{4})\s*:מתאריך")


class FamilyBizSchemaError(ValueError):
    """Raised when a workbook cannot be safely interpreted as FamilyBiz."""


@dataclass(frozen=True)
class ParsedWorkbook:
    inspection: ImportInspection
    records: list[ParsedSourceRecord]
    issues: list[DataQualityIssue]


def normalize_text(value: Any) -> str | None:
    if value is None or value == "":
        return None
    text = unicodedata.normalize("NFKC", str(value))
    text = text.replace("_x000D_", "\n")
    text = text.strip()
    while text.startswith("'"):
        text = text[1:].lstrip()
    return text.strip() or None


def normalize_decimal(value: Any) -> Decimal | None:
    if value is None or value == "":
        return None
    try:
        result = Decimal(str(value).strip().lstrip("'"))
    except (InvalidOperation, ValueError) as exc:
        raise FamilyBizSchemaError(f"Non-numeric amount: {value!r}") from exc
    if not result.is_finite():
        raise FamilyBizSchemaError(f"Non-finite amount: {value!r}")
    return result


def parse_date(value: Any, *, field: str, row_number: int) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = normalize_text(value)
    if not text:
        raise FamilyBizSchemaError(f"Missing {field} at source row {row_number}")
    try:
        return datetime.strptime(text, "%d/%m/%Y").date()  # noqa: DTZ007
    except ValueError as exc:
        raise FamilyBizSchemaError(
            f"Invalid {field} {text!r} at source row {row_number}; expected dd/mm/yyyy"
        ) from exc


def sha256_bytes(file_bytes: bytes) -> str:
    return hashlib.sha256(file_bytes).hexdigest()


def _safe_json_value(value: Any) -> Any:
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    return value


def _fingerprint(payload: Any) -> str:
    encoded = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(encoded.encode("utf-8")).hexdigest()


class FamilyBizParser:
    def __init__(self, settings: Settings | None = None) -> None:
        self.settings = settings or Settings.from_environment()

    def inspect(self, file_bytes: bytes, filename: str | None = None) -> ImportInspection:
        return self.parse(file_bytes, filename=filename).inspection

    def parse(self, file_bytes: bytes, filename: str | None = None) -> ParsedWorkbook:
        uncompressed_bytes, worksheet_max_row = self._validate_container(file_bytes, filename)
        file_hash = sha256_bytes(file_bytes)
        try:
            workbook = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=False)
        except (InvalidFileException, OSError, ValueError, zipfile.BadZipFile) as exc:
            raise FamilyBizSchemaError(f"Unable to read XLSX workbook: {exc}") from exc

        try:
            if len(workbook.sheetnames) != 1:
                raise FamilyBizSchemaError("Expected exactly one FamilyBiz worksheet")
            worksheet = workbook.active
            worksheet.reset_dimensions()
            records: list[ParsedSourceRecord] = []
            issues: list[DataQualityIssue] = []
            currencies: Counter[str] = Counter()
            account_kinds: Counter[str] = Counter()
            dates: list[date] = []
            preview_rows: list[dict[str, Any]] = []
            recent_rows: deque[tuple[int, list[Any]]] = deque(maxlen=3)
            report_start: date | None = None
            report_end: date | None = None
            active_account: AccountDescriptor | None = None
            section_count = 0
            pending_metadata_row: int | None = None

            for row_number, values in enumerate(
                worksheet.iter_rows(min_row=1, max_row=worksheet_max_row, max_col=10, values_only=True),
                start=1,
            ):
                raw_row = list(values)
                raw_row.extend([None] * (10 - len(raw_row)))
                if row_number == 1:
                    report_start = self._report_date(raw_row, REPORT_START_RE)
                    report_end = self._report_date(raw_row, REPORT_END_RE)

                header = self._is_header(raw_row, row_number)
                if header:
                    section_count += 1
                    # FamilyBiz places account metadata two worksheet rows before its header.
                    metadata_row_number = row_number - 2
                    prior = {number: row for number, row in recent_rows}
                    if metadata_row_number not in prior:
                        raise FamilyBizSchemaError(
                            f"Missing account identity before FamilyBiz section {section_count}"
                        )
                    if pending_metadata_row is not None and pending_metadata_row != metadata_row_number:
                        raise FamilyBizSchemaError(
                            f"Unexpected non-empty row at source row {pending_metadata_row}"
                        )
                    active_account = self._account_from_metadata(
                        prior[metadata_row_number], section_count
                    )
                    pending_metadata_row = None
                    recent_rows.append((row_number, raw_row))
                    continue

                if active_account is not None and any(value not in (None, "") for value in raw_row):
                    if pending_metadata_row is not None:
                        raise FamilyBizSchemaError(
                            f"Unexpected non-empty row at source row {row_number}"
                        )
                    if raw_row[0] not in (None, ""):
                        pending_metadata_row = row_number
                    elif raw_row[1] in (None, ""):
                        raise FamilyBizSchemaError(
                            f"Unexpected non-empty row without a date at source row {row_number}"
                        )
                    elif pending_metadata_row is None:
                        record = self._parse_record(
                            raw_row,
                            account=active_account,
                            source_row_number=row_number,
                            section_number=section_count,
                            sheet_name=worksheet.title,
                        )
                        records.append(record)
                        if len(records) > self.settings.max_rows:
                            raise FamilyBizSchemaError(
                                f"Workbook exceeds {self.settings.max_rows:,} transaction row limit"
                            )
                        issues.extend(record.issues)
                        currencies[record.currency] += 1
                        account_kinds[record.account.kind] += 1
                        dates.append(record.booking_date)
                        if len(preview_rows) < 20:
                            preview_rows.append(self._preview_row(record))
                elif pending_metadata_row is not None and row_number >= pending_metadata_row + 2:
                    raise FamilyBizSchemaError(
                        f"Unexpected non-empty row at source row {pending_metadata_row}"
                    )
                recent_rows.append((row_number, raw_row))

            if pending_metadata_row is not None:
                raise FamilyBizSchemaError(
                    f"Unexpected non-empty row at source row {pending_metadata_row}"
                )
            if not section_count:
                raise FamilyBizSchemaError("FamilyBiz headers not found")
            if not records:
                raise FamilyBizSchemaError("FamilyBiz worksheet contains no transaction rows")

            inspection = ImportInspection(
                parser_version=self.settings.parser_version,
                file_sha256=file_hash,
                filename=Path(filename).name if filename else None,
                compressed_bytes=len(file_bytes),
                uncompressed_bytes=uncompressed_bytes,
                sheet_names=list(workbook.sheetnames),
                section_count=section_count,
                transaction_count=len(records),
                report_start=report_start,
                report_end=report_end,
                min_booking_date=min(dates),
                max_booking_date=max(dates),
                currencies=dict(currencies),
                account_kinds=dict(account_kinds),
                issue_counts=dict(Counter(issue.code for issue in issues)),
                preview_rows=preview_rows,
            )
            return ParsedWorkbook(inspection=inspection, records=records, issues=issues)
        finally:
            workbook.close()

    def _validate_container(self, file_bytes: bytes, filename: str | None) -> tuple[int, int]:
        if not file_bytes:
            raise FamilyBizSchemaError("Empty upload")
        if filename and Path(filename).suffix.lower() not in {".xlsx"}:
            raise FamilyBizSchemaError("Only .xlsx FamilyBiz exports are supported")
        if len(file_bytes) > self.settings.max_compressed_bytes:
            raise FamilyBizSchemaError(
                "Compressed upload exceeds "
                f"{self.settings.max_compressed_bytes // 1024 // 1024} MiB limit"
            )
        try:
            with zipfile.ZipFile(io.BytesIO(file_bytes)) as archive:
                names = set(archive.namelist())
                if "xl/vbaProject.bin" in names or any(
                    name.endswith("vbaProject.bin") for name in names
                ):
                    raise FamilyBizSchemaError("Macro-enabled workbooks are not accepted")
                uncompressed = sum(info.file_size for info in archive.infolist())
                if uncompressed > self.settings.max_uncompressed_bytes:
                    raise FamilyBizSchemaError("Uncompressed upload exceeds 100 MiB limit")
                if archive.testzip() is not None:
                    raise FamilyBizSchemaError("Corrupt XLSX ZIP content")
                row_count = 0
                max_row = 0
                worksheet_entries = [
                    name
                    for name in archive.namelist()
                    if name.startswith("xl/worksheets/") and name.endswith(".xml")
                ]
                for worksheet_entry in worksheet_entries:
                    last_row_number = 0
                    try:
                        with archive.open(worksheet_entry) as worksheet_xml:
                            stack = []
                            for event, element in ElementTree.iterparse(
                                worksheet_xml, events=("start", "end")
                            ):
                                if event == "start":
                                    stack.append(element)
                                    continue
                                if element.tag.rsplit("}", 1)[-1] != "row":
                                    stack.pop()
                                    continue
                                row_count += 1
                                row_ref = element.attrib.get("r")
                                try:
                                    row_number = int(row_ref or "")
                                except ValueError as exc:
                                    raise FamilyBizSchemaError(
                                        "Worksheet row has an invalid XML coordinate"
                                    ) from exc
                                if row_number < 1 or row_number <= last_row_number:
                                    raise FamilyBizSchemaError(
                                        "Worksheet rows have invalid or duplicate XML coordinates"
                                    )
                                last_row_number = row_number
                                max_row = max(max_row, row_number)
                                if row_count > self.settings.max_rows + max(
                                    10, self.settings.max_rows // 20
                                ):
                                    raise FamilyBizSchemaError(
                                        f"Workbook exceeds {self.settings.max_rows:,} row limit"
                                    )
                                previous_column = 0
                                for cell in element:
                                    if cell.tag.rsplit("}", 1)[-1] != "c":
                                        continue
                                    coordinate = cell.attrib.get("r", "")
                                    match = re.fullmatch(
                                        r"([A-Z]+)([1-9][0-9]*)", coordinate
                                    )
                                    if not match or int(match.group(2)) != row_number:
                                        raise FamilyBizSchemaError(
                                            "Worksheet cell has an invalid XML coordinate"
                                        )
                                    try:
                                        column_number = column_index_from_string(match.group(1))
                                    except ValueError as exc:
                                        raise FamilyBizSchemaError(
                                            "Worksheet cell has an invalid XML coordinate"
                                        ) from exc
                                    if column_number <= previous_column:
                                        raise FamilyBizSchemaError(
                                            "Worksheet cells have invalid or duplicate XML coordinates"
                                        )
                                    previous_column = column_number
                                    has_value = any(
                                        child.tag.rsplit("}", 1)[-1] == "f"
                                        or (
                                            child.tag.rsplit("}", 1)[-1] in {"v", "is"}
                                            and any(text.strip() for text in child.itertext())
                                        )
                                        for child in cell
                                    )
                                    if has_value and column_number > 10:
                                        raise FamilyBizSchemaError(
                                            "Populated worksheet cell is outside the supported ten-column schema"
                                        )
                                if len(stack) > 1:
                                    stack[-2].remove(element)
                                element.clear()
                                stack.pop()
                    except ElementTree.ParseError as exc:
                        raise FamilyBizSchemaError("Malformed worksheet XML") from exc
                if max_row > self.settings.max_rows + max(10, self.settings.max_rows // 20):
                    raise FamilyBizSchemaError(
                        f"Workbook exceeds {self.settings.max_rows:,} row limit"
                    )
                return uncompressed, max_row
        except zipfile.BadZipFile as exc:
            raise FamilyBizSchemaError("Malformed XLSX ZIP content") from exc

    @staticmethod
    def _is_header(row: list[Any], source_row_number: int) -> bool:
        values = [normalize_text(value) for value in row[1:10]]
        if values == HEADERS:
            return True
        if values and values[0] == "תאריך" and any(values):
            raise FamilyBizSchemaError(
                f"Unknown or changed FamilyBiz headers at source row {source_row_number}"
            )
        return False

    @staticmethod
    def _report_date(row: list[Any], pattern: re.Pattern[str]) -> date | None:
        text = " ".join(normalize_text(value) or "" for value in row[:3])
        match = pattern.search(text)
        if not match:
            return None
        return datetime.strptime(match.group(1), "%d/%m/%Y").date()  # noqa: DTZ007

    @staticmethod
    def _account_from_metadata(
        metadata_row: list[Any], section_number: int
    ) -> AccountDescriptor:
        provider = normalize_text(metadata_row[0] if len(metadata_row) > 0 else None)
        reference = normalize_text(metadata_row[1] if len(metadata_row) > 1 else None)
        label = normalize_text(metadata_row[2] if len(metadata_row) > 2 else None)
        if not provider or not reference:
            raise FamilyBizSchemaError(
                f"Missing account identity before FamilyBiz section {section_number}"
            )
        kind = "bank" if "בנק" in provider else "card"
        currency = "USD" if reference.upper().endswith("USD") else "ILS"
        if label and label.upper() == "USD":
            currency = "USD"
        return AccountDescriptor(
            provider=provider,
            kind=kind,
            display_label=f"{provider} {kind}",
            source_reference_fingerprint=_fingerprint(
                {"provider": provider, "reference": reference}
            ),
            currency=currency,
        )

    def _parse_record(
        self,
        raw_row: list[Any],
        *,
        account: AccountDescriptor,
        source_row_number: int,
        section_number: int,
        sheet_name: str,
    ) -> ParsedSourceRecord:
        for index, value in enumerate(raw_row[1:10], start=1):
            if isinstance(value, str) and value.startswith("="):
                raise FamilyBizSchemaError(
                    "Formula in required FamilyBiz data cell at source row "
                    f"{source_row_number}, column {index + 1}"
                )
        booking_date = parse_date(
            raw_row[1], field="booking date", row_number=source_row_number
        )
        allocation_date = parse_date(
            raw_row[4], field="allocation date", row_number=source_row_number
        )
        amount = normalize_decimal(raw_row[2])
        if amount is None:
            raise FamilyBizSchemaError(f"Missing amount at source row {source_row_number}")
        description = normalize_text(raw_row[3])
        category = normalize_text(raw_row[6])
        currency = normalize_text(raw_row[7])
        if not description or not category or not currency:
            raise FamilyBizSchemaError(
                f"Missing required transaction value at source row {source_row_number}"
            )
        movement_type = normalize_text(raw_row[5])
        original_currency = normalize_text(raw_row[8])
        original_amount = normalize_decimal(raw_row[9])
        normalized_payload = [_safe_json_value(value) for value in raw_row]
        row_fingerprint = _fingerprint(
            {
                "account": account.source_reference_fingerprint,
                "row": normalized_payload,
            }
        )
        record_issues: list[DataQualityIssue] = []
        if not movement_type:
            record_issues.append(
                DataQualityIssue(
                    code="MISSING_MOVEMENT_TYPE",
                    message="Movement type is missing in the source row",
                    source_row=source_row_number,
                    section_index=section_number,
                )
            )
        if original_currency is None or original_amount is None:
            record_issues.append(
                DataQualityIssue(
                    code="MISSING_ORIGINAL_VALUE",
                    message="Original currency or amount is missing",
                    source_row=source_row_number,
                    section_index=section_number,
                )
            )
        if allocation_date != booking_date:
            record_issues.append(
                DataQualityIssue(
                    code="ALLOCATION_DATE_DIFFERS",
                    message="Allocation date differs from booking date",
                    source_row=source_row_number,
                    section_index=section_number,
                )
            )
        if amount == 0 and original_amount not in (None, Decimal(0)):
            record_issues.append(
                DataQualityIssue(
                    code="ZERO_AMOUNT_WITH_ORIGINAL_VALUE",
                    message="Normalized amount is zero while original amount is non-zero",
                    source_row=source_row_number,
                    section_index=section_number,
                )
            )
        if currency.upper() != "ILS":
            record_issues.append(
                DataQualityIssue(
                    code="NON_ILS_CURRENCY",
                    message="Non-ILS transaction is imported but excluded from ILS analytics",
                    source_row=source_row_number,
                    section_index=section_number,
                )
            )
        return ParsedSourceRecord(
            sheet_name=sheet_name,
            section_index=section_number,
            source_row_number=source_row_number,
            account=account,
            booking_date=booking_date,
            amount=amount,
            description=description,
            allocation_date=allocation_date,
            movement_type=movement_type,
            category=category,
            currency=currency,
            original_currency=original_currency,
            original_amount=original_amount,
            raw_payload=dict(zip(_RAW_PAYLOAD_KEYS, normalized_payload)),
            row_fingerprint=row_fingerprint,
            issues=record_issues,
        )

    @staticmethod
    def _preview_row(record: ParsedSourceRecord) -> dict[str, Any]:
        return {
            "section": record.section_index,
            "source_row": record.source_row_number,
            "account": record.account.display_label,
            "booking_date": record.booking_date.isoformat(),
            "amount": str(record.amount),
            "currency": record.currency,
            "issue_codes": [issue.code for issue in record.issues],
        }
