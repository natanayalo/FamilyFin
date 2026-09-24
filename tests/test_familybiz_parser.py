from __future__ import annotations

import pytest

from family_finance.importers.familybiz import FamilyBizParser, FamilyBizSchemaError

from .conftest import make_workbook


def test_parser_normalizes_rows_and_reports_warnings(familybiz_row):
    file_bytes = make_workbook(
        [
            familybiz_row,
            [
                "18/09/2026",
                0,
                "merchant 2",
                "19/09/2026",
                "",
                "category",
                "ILS",
                "USD",
                12,
            ],
        ]
    )
    parsed = FamilyBizParser().parse(file_bytes, filename="synthetic.xlsx")

    assert parsed.inspection.transaction_count == 2
    assert parsed.inspection.section_count == 1
    assert parsed.records[0].category == "household"
    assert parsed.records[1].movement_type is None
    assert {issue.code for issue in parsed.issues} >= {
        "MISSING_MOVEMENT_TYPE",
        "ALLOCATION_DATE_DIFFERS",
        "ZERO_AMOUNT_WITH_ORIGINAL_VALUE",
    }


def test_unknown_headers_fail_clearly(familybiz_row):
    file_bytes = bytearray(make_workbook([familybiz_row]))
    # Recreate the workbook with a changed header to avoid mutating ZIP XML manually.
    import io

    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(file_bytes))
    workbook.active.cell(5, 6).value = "unknown"
    output = io.BytesIO()
    workbook.save(output)

    with pytest.raises(FamilyBizSchemaError, match="headers"):
        FamilyBizParser().parse(output.getvalue(), filename="synthetic.xlsx")


def test_formula_in_required_data_cell_is_rejected(familybiz_row):
    row = list(familybiz_row)
    row[1] = "=1+1"
    with pytest.raises(FamilyBizSchemaError, match="Formula"):
        FamilyBizParser().parse(make_workbook([row]), filename="synthetic.xlsx")


def test_non_xlsx_extension_is_rejected(familybiz_row):
    with pytest.raises(FamilyBizSchemaError, match="Only .xlsx"):
        FamilyBizParser().parse(make_workbook([familybiz_row]), filename="synthetic.xlsm")


def test_corrupt_zip_is_rejected():
    with pytest.raises(FamilyBizSchemaError, match="Malformed XLSX"):
        FamilyBizParser().parse(b"not an xlsx", filename="synthetic.xlsx")


def test_unexpected_nonempty_row_is_rejected(familybiz_row):
    from io import BytesIO

    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(make_workbook([familybiz_row])))
    workbook.active.cell(7, 1).value = "unexpected"
    workbook.active.cell(7, 2).value = "not-account-metadata"
    output = BytesIO()
    workbook.save(output)

    with pytest.raises(FamilyBizSchemaError, match="Unexpected non-empty row"):
        FamilyBizParser().parse(output.getvalue(), filename="synthetic.xlsx")


def test_transaction_row_limit_is_enforced(familybiz_row):
    from family_finance.config import Settings

    settings = Settings(max_rows=0)
    with pytest.raises(FamilyBizSchemaError, match="transaction row limit"):
        FamilyBizParser(settings).parse(make_workbook([familybiz_row]), filename="synthetic.xlsx")


def test_populated_cells_outside_supported_schema_are_rejected(familybiz_row):
    from io import BytesIO

    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(make_workbook([familybiz_row])))
    workbook.active.cell(6, 11).value = "outside schema"
    output = BytesIO()
    workbook.save(output)

    with pytest.raises(FamilyBizSchemaError, match="outside the supported ten-column"):
        FamilyBizParser().parse(output.getvalue(), filename="synthetic.xlsx")


def test_actual_xml_cell_row_coordinate_is_validated(familybiz_row):
    from io import BytesIO
    from zipfile import ZIP_DEFLATED, ZipFile

    source = make_workbook([familybiz_row])
    rewritten = BytesIO()
    with ZipFile(BytesIO(source)) as original, ZipFile(rewritten, "w", ZIP_DEFLATED) as target:
        for item in original.infolist():
            payload = original.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                payload = payload.replace(b'r="B6"', b'r="B7"', 1)
            target.writestr(item, payload)

    with pytest.raises(FamilyBizSchemaError, match="cell has an invalid XML coordinate"):
        FamilyBizParser().parse(rewritten.getvalue(), filename="synthetic.xlsx")
