"""Destructive only to explicitly marked disposable labs; never a live project.

Run once against fresh lab volumes. Credentials exist only in process memory and
are sent to docker exec over stdin. HTTP probes simulate the TLS proxy boundary;
they do not certify a browser or a phone's HTTPS connection.
"""
import json
import secrets
import socket
from datetime import UTC, datetime

from lab_profile import require_lab_python
from manage import (
    PROJECTS,
    REPORTS,
    STATE,
    assert_owned,
    baseline,
    baseline_equal,
    compose,
    engine,
    inspect_all,
    request,
    run,
    verify,
)

require_lab_python()

PRIMARY = "familyfin-m1-synthetic"
RESTORED = "familyfin-m1-restore"
results = []


def record(name, function):
    try:
        detail = function()
        results.append({"test": name, "status": "PASS", "evidence": detail or "Executed assertions passed"})
        print("PASS " + name, flush=True)
    except Exception as exc:
        results.append({"test": name, "status": "FAIL", "evidence": str(exc)})
        raise
    finally:
        REPORTS.mkdir(exist_ok=True)
        (REPORTS / "synthetic-validation.json").write_text(json.dumps({
            "checked_at": datetime.now(UTC).isoformat(), "results": results,
            "scope": "Synthetic localhost production stack; no Serve modification or mobile acceptance"
        }, indent=2), encoding="utf-8")


def execute(project, code, data=None):
    return compose(project, "exec", "-T", "api", "python", "-c", code,
                   stdin=json.dumps(data) if data is not None else None).stdout.strip()


def status(project):
    return json.loads(execute(project, '''
import hashlib,json,sqlite3
from pathlib import Path
root=Path('/var/lib/familyfin')
with sqlite3.connect(root/'family_finance.sqlite3') as c:
    tables={r[0] for r in c.execute("SELECT name FROM sqlite_master WHERE type='table'")}
    counts={t:c.execute('SELECT count(*) FROM '+t).fetchone()[0] for t in
            ['api_users','actor_audit_events','transactions','source_files','net_worth_accounts'] if t in tables}
    migrations=c.execute('SELECT version_num FROM alembic_version').fetchall()
    # Compare financial values, identities and revisions, not only row counts.
    # Archive path relocation is intentional; files are compared by content hash.
    financial={}
    for table in sorted(tables):
        if table.startswith(('api_', 'sqlite_', 'actor_')) or table == 'alembic_version':
            continue
        columns=[r[1] for r in c.execute('PRAGMA table_info("'+table+'")')
                 if r[1] not in ('archived_path','managed_path')]
        rows=c.execute('SELECT '+','.join('"'+col+'"' for col in columns)+' FROM "'+table+'"').fetchall()
        financial[table]={'columns':columns,'rows':sorted(rows,key=repr)}
    financial_hash=hashlib.sha256(json.dumps(financial,sort_keys=True,default=str).encode()).hexdigest()
archives={str(p.relative_to(root)):hashlib.sha256(p.read_bytes()).hexdigest()
          for folder in ['imports','planning-imports','net-worth-imports']
          for p in (root/folder).rglob('*') if p.is_file()}
print(json.dumps({'counts':counts,'migrations':migrations,'archives':archives,
                 'financial_sha256':financial_hash,
                 'key':hashlib.sha256((root/'api-secret.key').read_bytes()).hexdigest()}))
'''))


def login(project, username, password):
    status_code, headers, body = request(PROJECTS[project], "/api/v1/auth/session", "POST",
                                        data={"username": username, "password": password})
    assert status_code == 200, f"Login failed: {status_code}"
    cookie = headers["set-cookie"]
    for flag in ("secure", "httponly", "samesite=strict", "path=/api/v1"):
        assert flag in cookie.lower()
    payload = json.loads(body)["data"]
    assert payload["user"]["username"] == username
    return cookie.split(";", 1)[0], payload["csrf_token"]


def main():
    engine()
    for project in PROJECTS:
        assert_owned(project)
        existing = run("docker", "volume", "inspect", project + "_financial", check=False)
        if existing.returncode == 0:
            raise SystemExit("Fresh lab volumes required. Preserve/review previous evidence; do not run against existing data.")
    STATE.mkdir(exist_ok=True)
    before = baseline()
    (STATE / "baseline.json").write_text(json.dumps(before, indent=2), encoding="utf-8")
    for port in PROJECTS.values():
        with socket.socket() as sock:
            sock.bind(("127.0.0.1", port))

    def start():
        compose(PRIMARY, "config", "--quiet")
        compose(PRIMARY, "up", "-d", "--wait", "--wait-timeout", "180")
        return verify(PRIMARY)
    record("01 production stack, migrations, health and hardening", start)

    passwords = [secrets.token_urlsafe(32), secrets.token_urlsafe(32)]

    def bootstrap():
        execute(PRIMARY, '''
import json,sys,os,io
from pathlib import Path
from openpyxl import Workbook
from family_finance.api.auth import AuthService,BootstrapAccount
from family_finance.config import Settings
from family_finance.persistence.db import Database
from family_finance.services import ImportService
from family_finance.importers.familybiz import HEADERS
os.umask(0o077)
s=Settings.from_environment(); d=Database(s.database_path)
assert not list(s.archive_root.iterdir())
AuthService(d,s).bootstrap_accounts([BootstrapAccount('qa-one','Synthetic One',p) if i==0
    else BootstrapAccount('qa-two','Synthetic Two',p) for i,p in enumerate(json.load(sys.stdin))])
(s.data_root/'.synthetic-lab').write_text('Disposable FamilyFin M1 validation only')
w=Workbook(); sh=w.active
sh.cell(1,1).value='30/09/2026 :עד תאריך'; sh.cell(1,2).value='01/09/2025 :מתאריך '
sh.cell(3,1).value='כאל'; sh.cell(3,2).value='synthetic-masked'
for i,v in enumerate(HEADERS,2): sh.cell(5,i).value=v
row=['19/09/2026',-17.4,'Synthetic merchant','19/09/2026','food','household','ILS','ILS',-17.4]
for i,v in enumerate(row,2): sh.cell(6,i).value=v
out=io.BytesIO(); w.save(out); payload=out.getvalue()
service=ImportService(settings=s,database=d); preview=service.preview_import(payload)
service.commit_import(payload,preview.preview_token,'synthetic.xlsx')
assert (s.data_root/'api-secret.key').stat().st_mode & 0o777 == 0o600
assert s.database_path.stat().st_mode & 0o077 == 0
assert s.data_root.stat().st_uid == 10001 and s.data_root.stat().st_mode & 0o777 == 0o700
''', passwords)
        current = status(PRIMARY)
        assert current["counts"]["api_users"] == 2
        assert current["counts"]["transactions"] == 1 and current["archives"]
        return "Two random-password accounts; owner-only key/database; synthetic XLSX and transaction archived"
    record("02 secure synthetic bootstrap and source archive", bootstrap)
    sessions = []

    def auth_checks():
        port = PROJECTS[PRIMARY]
        for path in ("/api/v1/auth/session", "/api/v1/net-worth/accounts", "/api/v1/planning/scenarios", "/api/v1/automation/runs"):
            assert request(port, path)[0] == 401
        for origin in (None, "https://attacker.example"):
            assert request(port, "/api/v1/auth/session", "POST", data={"username":"qa-one", "password":passwords[0]}, origin=origin)[0] == 403
        assert request(port, "/api/v1/auth/session", "POST", data={"username":"qa-one", "password":"incorrect"})[0] == 401
        for i, password in enumerate(passwords):
            username = ["qa-one", "qa-two"][i]
            cookie, csrf = login(PRIMARY, username, password)
            assert request(port, "/api/v1/auth/session", cookie=cookie)[0] == 200
            body = {"account_key": f"synthetic-{i}", "display_name": f"Synthetic {i}",
                        "side": "asset", "category": "cash", "liquidity": "liquid", "active_from": "2026-09-01"}
            assert request(port, "/api/v1/net-worth/accounts", "POST", data=body, cookie=cookie)[0] == 403
            assert request(port, "/api/v1/net-worth/accounts", "POST", data=body, cookie=cookie, csrf="wrong")[0] == 403
            assert request(port, "/api/v1/net-worth/accounts", "POST", data=body, cookie=cookie, csrf=csrf, origin="https://attacker.example")[0] == 403
            assert request(port, "/api/v1/net-worth/accounts", "POST", data=body, cookie=cookie, csrf=csrf)[0] == 201
            assert request(port, "/api/v1/auth/session", "DELETE", cookie=cookie, csrf=csrf)[0] == 204
            assert request(port, "/api/v1/auth/session", cookie=cookie)[0] == 401
            sessions.append(login(PRIMARY, username, password))
        assert request(port, "/api/v1/auth/session", cookie=sessions[0][0], csrf=sessions[1][1], method="DELETE")[0] == 403
        assert request(port, "/api/v1/auth/session", cookie="familyfin_session=invalid")[0] == 401
        assert request(port, "/api/v1/auth/session", "POST", raw=b"x"*(1024*1024+1))[0] == 413
        return "Both users: login, secure cookie attributes, equal write permissions, missing/wrong/cross-session CSRF, origin, logout/revocation; auth boundaries and JSON limit"
    record("03 authentication, sessions, CSRF and request limits", auth_checks)

    initial = status(PRIMARY)

    def persistence():
        compose(PRIMARY, "up", "-d", "--force-recreate", "--wait", "--wait-timeout", "180")
        assert status(PRIMARY) == initial
        assert request(PROJECTS[PRIMARY], "/api/v1/auth/session", cookie=sessions[0][0])[0] == 200
        return verify(PRIMARY)
    record("04 persistent data and sessions after container recreation", persistence)

    def isolation():
        assert execute(PRIMARY, "from pathlib import Path; assert not Path('/app/src').stat().st_mode & 0o002; print('ok')") == "ok"
        for path in ("/app/should-not-write", "/etc/should-not-write"):
            result = compose(PRIMARY, "exec", "-T", "api", "python", "-c", f"open({path!r},'w').close()", check=False)
            assert result.returncode != 0
        compose(PRIMARY, "exec", "-T", "web", "node", "-e",
                "const f=require('fs');if(f.existsSync('/var/lib/familyfin')||f.existsSync('/backups')||f.existsSync('/var/run/docker.sock'))process.exit(1)")
        # No financial mount exists in any unrelated container; no commands are
        # executed inside those containers, nor are their files read.
        verify(PRIMARY)
        return "API rootfs write denied; web sees no financial/backup path or socket; daemon mount inventory confirms exclusivity"
    record("05 filesystem and unrelated-container volume isolation", isolation)

    snapshot = []

    def backup():
        name = compose(PRIMARY, "run", "--rm", "--no-deps", "operations", "operations", "backup").stdout.strip().splitlines()[-1]
        assert name.startswith("familyfin-") and "/" not in name
        compose(PRIMARY, "run", "--rm", "--no-deps", "operations", "operations", "verify", name)
        snapshot.append(name)
        assert status(PRIMARY) == initial
        return "Read-only audit, SQLite online snapshot, managed source archive, manifest/hash/schema verification; live state unchanged"
    record("06 full application backup and independent verification", backup)

    def restore():
        source = PRIMARY + "_backups"
        compose(RESTORED, "run", "--rm", "--no-deps", "operations", "operations", "restore", snapshot[0], restore_source=source)
        compose(RESTORED, "up", "-d", "--wait", "--wait-timeout", "180")
        restored = status(RESTORED)
        assert restored["key"] != initial["key"]
        assert {k:v for k,v in restored.items() if k != "key"} == {k:v for k,v in initial.items() if k != "key"}
        for cookie, _ in sessions:
            assert request(PROJECTS[RESTORED], "/api/v1/auth/session", cookie=cookie)[0] == 401
        for i, password in enumerate(passwords):
            cookie, _csrf = login(RESTORED, ["qa-one", "qa-two"][i], password)
            code, _, body = request(PROJECTS[RESTORED], "/api/v1/net-worth/accounts", cookie=cookie)
            assert code == 200 and len(json.loads(body)["data"]) == 2
        execute(RESTORED, "from family_finance.audit import AuditService; assert AuditService().run().passed")
        # Fail-closed protection for a nonempty recovery root.
        rejected = compose(RESTORED, "run", "--rm", "--no-deps", "operations", "operations", "restore", snapshot[0], restore_source=source, check=False)
        assert rejected.returncode != 0 and "empty" in rejected.stderr
        fresh = compose(RESTORED, "run", "--rm", "--no-deps", "operations", "operations", "backup").stdout.strip().splitlines()[-1]
        compose(RESTORED, "run", "--rm", "--no-deps", "operations", "operations", "verify", fresh)
        return "Separate clean project/volume, matching financial rows+archive hashes+migration+actor counts, old sessions rejected, both users reauthenticate, audit and fresh verified backup pass; nonempty restore rejected"
    record("07 clean restore, financial state, audit and reauthentication", restore)

    def ports():
        verify(PRIMARY); verify(RESTORED)
        # Probe own current LAN address. Independent peer and IPv6 tests remain
        # explicitly UNVERIFIED in the acceptance report.
        ps = run("powershell", "-NoProfile", "-Command",
                 "Get-NetIPAddress -AddressFamily IPv4 | Where-Object { $_.InterfaceAlias -eq 'Wi-Fi' -and $_.IPAddress -notlike '169.254.*' } | Select-Object -ExpandProperty IPAddress").stdout.split()
        assert ps, "No LAN address discovered"
        for address in ps:
            for port in (*PROJECTS.values(), 8000, 3000):
                with socket.socket() as sock:
                    sock.settimeout(2)
                    assert sock.connect_ex((address, port)) != 0, "Unexpected LAN response"
        for item in inspect_all():
            bindings = item["HostConfig"].get("PortBindings") or {}
            for values in bindings.values():
                for binding in values or []:
                    assert binding["HostIp"] == "127.0.0.1", "Non-loopback Docker binding"
        assert baseline_equal(baseline(), before), "Existing service or Serve change"
        return f"Both lab ports loopback-only; all Docker bindings enumerated; own-host LAN probes to {list(PROJECTS.values())} and 8000/3000 denied; pre-existing IDs/configuration/runtime records and Serve identical"
    record("08 host ports, LAN probes and existing service preservation", ports)
    results.extend([
        {"test": "09 real browser service-worker/offline restrictions", "status": "UNVERIFIED", "evidence": "Worker handler harness is separate evidence; phone cache and install acceptance requires approved HTTPS cutover"},
        {"test": "10 independent LAN peer, tailnet negative device, Wi-Fi/cellular mobile HTTPS", "status": "UNVERIFIED", "evidence": "No M1 Serve route activated; M0 observations cannot certify M1"},
        {"test": "11 encrypted independent backup media and scheduled recovery", "status": "UNVERIFIED", "evidence": "Lab uses a distinct Docker volume on the same virtual disk; production encrypted destination and scheduler not approved/configured"},
    ])
    (REPORTS / "synthetic-validation.json").write_text(json.dumps({"checked_at":datetime.now(UTC).isoformat(), "results":results}, indent=2), encoding="utf-8")
    print(f"Synthetic stack and restore lab on localhost {list(PROJECTS.values())}. No live endpoint activated.")


if __name__ == "__main__":
    main()
